import importlib.util,tempfile,unittest
from pathlib import Path
from decimal import Decimal
from datetime import date
import openpyxl
spec=importlib.util.spec_from_file_location('legacy',Path(__file__).with_name('prepare-legacy-import.py'));m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
class ConverterTests(unittest.TestCase):
 def test_residual_reconciles_exactly(self):
  values,residual=m.reconcile([Decimal('1.00006'),Decimal('2.00006')],Decimal('3.0001'))
  self.assertEqual(sum(values),Decimal('3.0001'));self.assertEqual(residual,'-0.0001')
 def test_large_discrepancy_rejected(self):
  with self.assertRaises(m.ValidationError):m.reconcile([Decimal('1')],Decimal('2'))
 def test_private_conversion_and_failure_writes_nothing(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);book=openpyxl.Workbook();s=book.active
   for c,v in {2:'NO',3:'NAME',4:'PRINCIPAL',7:'PRINCIPAL',10:'PRINCIPAL'}.items():s.cell(3,c,v)
   s.cell(4,2,1);s.cell(4,3,'Anonymous fixture')
   values=[10,2,12,5,1,6,15,3,18]
   for c,v in enumerate(values,4):s.cell(4,c,v)
   s.cell(5,3,'TOTAL')
   for c,v in enumerate(values,4):s.cell(5,c,v)
   source=root/'fixture.xlsx';book.save(source)
   report=m.prepare(source,root/'ok',date(2026,10,3));self.assertEqual(report['closing_equity'],'18.0000');self.assertNotIn('Anonymous',str(report))
   self.assertIn('SOURCE_BASELINE',(root/'ok'/'members.csv').read_text())
   s.cell(5,4,999);book.save(source)
   with self.assertRaises(m.ValidationError):m.prepare(source,root/'bad',date(2026,10,3))
   self.assertFalse((root/'bad').exists())
if __name__=='__main__':unittest.main()
